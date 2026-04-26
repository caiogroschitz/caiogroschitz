"""Generate a structured .docx report from validated jurisprudence items."""

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from ..models.jurisprudencia import Jurisprudencia, Favorabilidade, ResultadoColeta


def _slug(text: str) -> str:
    """Convert tema to a safe filesystem slug."""
    slug = re.sub(r"[^\w\s-]", "", text.lower())
    slug = re.sub(r"[\s-]+", "_", slug)
    return slug[:80].strip("_")


def _set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _cor_favorabilidade(fav: Favorabilidade) -> tuple[str, str]:
    """Return (hex_bg, label) for a favourability value."""
    return {
        Favorabilidade.FAVORAVEL: ("C6EFCE", "FAVORÁVEL"),
        Favorabilidade.DESFAVORAVEL: ("FFC7CE", "DESFAVORÁVEL"),
        Favorabilidade.NEUTRA: ("FFEB9C", "NEUTRA"),
        Favorabilidade.NAO_CLASSIFICADA: ("FFFFFF", "NÃO CLASSIFICADA"),
    }[fav]


def gerar_docx(resultado: ResultadoColeta) -> str:
    """Write a .docx to disk and return the absolute file path."""
    pasta = Path.home() / "Jurisprudencias" / _slug(resultado.tema)
    pasta.mkdir(parents=True, exist_ok=True)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    nome_arquivo = f"jurisprudencias_{_slug(resultado.tema)}_{ts}.docx"
    caminho = pasta / nome_arquivo

    doc = Document()

    # ---------- Page margins ----------
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    # ---------- Title ----------
    titulo = doc.add_heading("Relatório de Jurisprudência", level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Tema: {resultado.tema}")
    doc.add_paragraph(
        f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    )
    doc.add_paragraph(
        f"Total coletado: {resultado.total_valido} | "
        f"Descartados: {resultado.total_descartado}"
    )
    doc.add_paragraph("")

    # ---------- Decisions ----------
    for idx, item in enumerate(resultado.itens, start=1):
        cor_bg, rotulo_fav = _cor_favorabilidade(item.favorabilidade)

        # Section heading
        heading = doc.add_heading(
            f"{idx}. {item.tribunal} — {item.numero_processo}",
            level=2,
        )

        # Metadata table
        table = doc.add_table(rows=5, cols=2)
        table.style = "Table Grid"

        labels = [
            ("Tribunal", item.tribunal),
            ("Número do Processo", item.numero_processo),
            ("Data de Julgamento", item.data_julgamento.strftime("%d/%m/%Y")),
            ("Tipo de Decisão", item.tipo.value),
            ("Favorabilidade", rotulo_fav),
        ]
        for i, (label, value) in enumerate(labels):
            row = table.rows[i]
            row.cells[0].text = label
            run = row.cells[0].paragraphs[0].runs[0]
            run.bold = True
            run.font.size = Pt(10)
            row.cells[1].text = value
            row.cells[1].paragraphs[0].runs[0].font.size = Pt(10)

        # Colorize favorabilidade row
        _set_cell_bg(table.rows[4].cells[0], cor_bg)
        _set_cell_bg(table.rows[4].cells[1], cor_bg)

        # Link
        p_link = doc.add_paragraph()
        p_link.add_run("Link: ").bold = True
        p_link.add_run(item.link).font.color.rgb = RGBColor(0x00, 0x00, 0xFF)

        # Ementa — literal, verbatim
        doc.add_paragraph("Ementa:").runs[0].bold = True
        ementa_p = doc.add_paragraph(item.ementa)
        ementa_p.paragraph_format.left_indent = Inches(0.4)
        for run in ementa_p.runs:
            run.font.size = Pt(10)

        if item.relator:
            doc.add_paragraph(f"Relator: {item.relator}")

        doc.add_paragraph("")  # spacer

    # ---------- Discards summary ----------
    if resultado.descartados:
        doc.add_heading("Itens Descartados", level=2)
        t = doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        hdr = t.rows[0].cells
        hdr[0].text = "Identificador"
        hdr[1].text = "Motivo"
        for h in hdr:
            h.paragraphs[0].runs[0].bold = True
        for d in resultado.descartados:
            row = t.add_row().cells
            row[0].text = str(d.get("id", ""))
            row[1].text = str(d.get("motivo", ""))

    doc.save(str(caminho))
    return str(caminho)
