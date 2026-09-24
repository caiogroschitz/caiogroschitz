#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_roteiro_docx.py — Gera o ROTEIRO DE PREENCHIMENTO do PJe-Calc em .docx
a partir de um JSON estruturado.

Por que um script: o conteúdo do roteiro varia muito de processo para processo,
mas a FORMATAÇÃO (ordem das telas do PJe-Calc, citação de fonte, marcação de
itens a conferir) deve ser sempre consistente e profissional. O script garante
essa consistência; o JSON carrega o conteúdo extraído dos autos.

USO:
    python3 build_roteiro_docx.py <entrada.json> <saida.docx>

ESQUEMA DO JSON (todos os campos de texto aceitam string; listas podem ser
vazias). Campos marcados (opcional) podem ser omitidos.

{
  "processo": {
    "numero": "0001234-56.2024.5.02.0001",
    "vara": "1ª Vara do Trabalho de São Paulo/SP (TRT2)",
    "reclamante": "Fulano de Tal",
    "reclamado": "Empresa X Ltda",
    "titulo_fonte": "Sentença ID 1a2b3c, fls. 210/216; trânsito em julgado fl. 240"
  },
  "lembretes_iniciais": [
    "Atualizar as tabelas/índices on-line no PJe-Calc antes de iniciar."
  ],
  "parametros_processo": [
    {"campo": "Data de admissão", "valor": "09/06/2014",
     "fonte": "CTPS fl. 33", "observacao": ""}
  ],
  "historico_salarial": [
    {"periodo": "07/2014 a 04/2015", "salario": "R$ 935,00",
     "fonte": "Holerites fls. 50/61", "observacao": ""}
  ],
  "faltas_ferias": [
    {"tipo": "Falta injustificada", "periodo": "24 a 26/09/2014",
     "fonte": "Cartão de ponto fl. 70", "observacao": ""}
  ],
  "verbas": [
    {
      "nome": "Horas extras além da 8ª diária, adicional 50%",
      "periodo": "01/2015 a 09/2016",
      "base_calculo": "salário base + adicional de periculosidade",
      "adicional_divisor": "50% / divisor 220",
      "natureza": "Salarial",
      "incidencias": ["FGTS", "INSS", "IRPF"],
      "reflexos": ["Aviso prévio indenizado", "13º", "Férias + 1/3", "DSR", "FGTS + 40%"],
      "deducoes": "Descontar HE já pagas (holerites fls. 50/61)",
      "fonte": "Sentença item 'a', fl. 212",
      "conferir": ""
    }
  ],
  "correcao_juros": {
    "regime": "Conforme título: ... / ou regime vigente das tabelas do PJe-Calc",
    "marco_correcao": "Vencimento de cada parcela (época própria)",
    "juros_data_inicial": "Ajuizamento (10/04/2017)",
    "fonte": "Sentença fl. 215",
    "conferir": "Período cruza 30/08/2024 (Lei 14.905/2024) — validar na tabela."
  },
  "itens_conferir": [
    "Salário de 05/2016 diverge entre holerite e CTPS — segui a sentença (fl. 212)."
  ],
  "encerramento": [
    "Rodar a checagem de inconsistências do PJe-Calc.",
    "Conferir o relatório contra o dispositivo da sentença.",
    "Exportar .pjc e PDF; juntar ambos na petição (TRT2)."
  ]
}

Qualquer campo "conferir"/"itens_conferir" não vazio é renderizado com destaque,
para o usuário revisar antes de lançar.
"""

import json
import sys
from datetime import date

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    sys.exit("python-docx não encontrado. Instale: pip install python-docx --break-system-packages")

WARN = RGBColor(0xB0, 0x00, 0x00)      # vermelho para "⚠ Conferir"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)    # azul-petróleo para títulos
GREY = RGBColor(0x55, 0x55, 0x55)


def _set_base_style(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)


def _h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = ACCENT
    return p


def _source(p, fonte):
    if fonte:
        r = p.add_run(f"  [fonte: {fonte}]")
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GREY


def _warn_para(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("⚠ Conferir: ")
    r.bold = True
    r.font.color.rgb = WARN
    r2 = p.add_run(text)
    r2.font.color.rgb = WARN
    return p


def _kv_table(doc, rows, headers):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for par in hdr[i].paragraphs:
            for run in par.runs:
                run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val if val else ""
    return t


def build(data, out_path):
    doc = Document()
    _set_base_style(doc)

    # Cabeçalho
    title = doc.add_heading("Roteiro de preenchimento do PJe-Calc — Liquidação", level=0)
    for run in title.runs:
        run.font.color.rgb = ACCENT
    proc = data.get("processo", {})
    sub = doc.add_paragraph()
    sub.add_run("Parte autora / liquidação de sentença\n").bold = True
    for label, key in [("Processo", "numero"), ("Vara/Órgão", "vara"),
                       ("Reclamante", "reclamante"), ("Reclamado", "reclamado")]:
        if proc.get(key):
            sub.add_run(f"{label}: ").bold = True
            sub.add_run(f"{proc.get(key)}\n")
    if proc.get("titulo_fonte"):
        r = sub.add_run(f"Título executivo: {proc.get('titulo_fonte')}\n")
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
    g = doc.add_paragraph()
    gr = g.add_run(f"Gerado em {date.today().strftime('%d/%m/%Y')}. "
                   "Documento de apoio: a engine oficial é o PJe-Calc; confira o "
                   "cálculo final antes da juntada.")
    gr.italic = True
    gr.font.size = Pt(9)
    gr.font.color.rgb = GREY

    # Lembretes iniciais
    lembretes = data.get("lembretes_iniciais") or []
    if lembretes:
        _h(doc, "0. Antes de começar", level=1)
        for item in lembretes:
            doc.add_paragraph(item, style="List Bullet")

    # 1. Dados do Processo / Parâmetros
    params = data.get("parametros_processo") or []
    if params:
        _h(doc, "1. Dados do Processo e Parâmetros do Cálculo", level=1)
        doc.add_paragraph("Lance estes valores nas abas de identificação e de "
                          "parâmetros do PJe-Calc:")
        rows = [[p.get("campo", ""), p.get("valor", ""), p.get("fonte", ""),
                 p.get("observacao", "")] for p in params]
        _kv_table(doc, rows, ["Campo do PJe-Calc", "Valor a lançar", "Fonte (ID/fl.)", "Obs."])
        for p in params:
            if p.get("conferir"):
                _warn_para(doc, f"{p.get('campo','')}: {p['conferir']}")

    # 2. Histórico Salarial
    hist = data.get("historico_salarial") or []
    if hist:
        _h(doc, "2. Histórico Salarial", level=1)
        doc.add_paragraph("Aba 'Histórico Salarial' — informe por período:")
        rows = [[h.get("periodo", ""), h.get("salario", ""), h.get("fonte", ""),
                 h.get("observacao", "")] for h in hist]
        _kv_table(doc, rows, ["Período", "Salário", "Fonte (ID/fl.)", "Obs."])

    # 3. Faltas e Férias
    ff = data.get("faltas_ferias") or []
    if ff:
        _h(doc, "3. Faltas e Férias", level=1)
        rows = [[f.get("tipo", ""), f.get("periodo", ""), f.get("fonte", ""),
                 f.get("observacao", "")] for f in ff]
        _kv_table(doc, rows, ["Tipo", "Período/Data", "Fonte (ID/fl.)", "Obs."])

    # 4. Parcelas (verbas)
    verbas = data.get("verbas") or []
    if verbas:
        _h(doc, "4. Parcelas (verbas) — Principais e Reflexas", level=1)
        doc.add_paragraph("Lance cada verba como parcela no PJe-Calc, marcando "
                          "natureza, incidências e reflexos exatamente como abaixo "
                          "(reproduzindo o comando da sentença):")
        for idx, v in enumerate(verbas, 1):
            ph = doc.add_paragraph()
            ph.add_run(f"4.{idx} {v.get('nome','(verba)')}").bold = True
            _source(ph, v.get("fonte", ""))

            def line(label, value):
                if value:
                    pp = doc.add_paragraph(style="List Bullet")
                    pp.add_run(f"{label}: ").bold = True
                    pp.add_run(value)

            line("Período de incidência", v.get("periodo", ""))
            line("Base de cálculo", v.get("base_calculo", ""))
            line("Adicional / divisor", v.get("adicional_divisor", ""))
            line("Natureza", v.get("natureza", ""))
            inc = v.get("incidencias") or []
            if inc:
                line("Incidências (marcar)", ", ".join(inc))
            refl = v.get("reflexos") or []
            if refl:
                line("Reflexos (repercussões)", ", ".join(refl))
            line("Deduções", v.get("deducoes", ""))
            if v.get("conferir"):
                _warn_para(doc, v["conferir"])

    # 5. Correção, Juros e Multa
    cj = data.get("correcao_juros") or {}
    if cj:
        _h(doc, "5. Correção, Juros e Multa", level=1)
        for label, key in [("Regime de correção/juros", "regime"),
                           ("Marco inicial da correção", "marco_correcao"),
                           ("Juros — data inicial", "juros_data_inicial")]:
            if cj.get(key):
                pp = doc.add_paragraph(style="List Bullet")
                pp.add_run(f"{label}: ").bold = True
                pp.add_run(cj[key])
        if cj.get("fonte"):
            pf = doc.add_paragraph()
            _source(pf, cj["fonte"])
        if cj.get("conferir"):
            _warn_para(doc, cj["conferir"])

    # 6. Itens a conferir (consolidado)
    itens = data.get("itens_conferir") or []
    if itens:
        _h(doc, "6. Pontos de atenção antes de lançar", level=1)
        for it in itens:
            _warn_para(doc, it)

    # 7. Encerramento
    enc = data.get("encerramento") or []
    if enc:
        _h(doc, "7. Finalização no PJe-Calc", level=1)
        for it in enc:
            doc.add_paragraph(it, style="List Number")

    doc.save(out_path)
    return out_path


def main():
    if len(sys.argv) != 3:
        sys.exit("Uso: python3 build_roteiro_docx.py <entrada.json> <saida.docx>")
    in_path, out_path = sys.argv[1], sys.argv[2]
    with open(in_path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    path = build(data, out_path)
    print(f"Roteiro gerado: {path}")


if __name__ == "__main__":
    main()
