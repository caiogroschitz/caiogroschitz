#!/usr/bin/env python3
"""
Converte a peça de apelação escrita em Markdown para .docx no padrão BFAP.

Uso:
    python3 scripts/montar_docx.py --in peca.md --out "Apelacao - 0811684-11.docx"

O template `assets/modelo-bfap.docx` carrega o timbre, o rodapé, as margens e a
fonte Prompt (embutida) do escritório. O script só despeja o texto dentro dele,
então a peça sai com a identidade visual correta em qualquer máquina.

Marcação aceita no arquivo de entrada
-------------------------------------
    # TEXTO               título centralizado, negrito, sublinhado (RAZÕES DE APELAÇÃO)
    ## TEXTO              título de seção, negrito, justificado (I - DA TEMPESTIVIDADE)
    ^ TEXTO               parágrafo centralizado (endereçamento, fecho, assinatura)
    > texto               citação de lei/sentença: recuo de 2 cm, 11 pt
    >> texto              ementa de acórdão: recuo de 2 cm, 10 pt
    [[PROVA 3 | print do cadastro do apelado desde 24/08/2017]]
                          moldura vazia para o advogado colar o print
    [[QUADRO]]            quadro de identificação das partes (uma linha por par)
    APELANTE: ...
    APELADO: ...
    [[/QUADRO]]
    ---                   quebra de página (separa a interposição das razões)
    **negrito**  *itálico*  dentro de qualquer parágrafo

Qualquer outra linha vira parágrafo comum justificado. Linhas em branco são ignoradas.
"""

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Prompt"
BODY_PT = 12
CITE_PT = 11
EMENTA_PT = 10
LINE = 1.15
SPACE = Pt(10)
INDENT = Cm(2)


def _style_run(run, size, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    # garante a fonte também para East Asian / Complex Script
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)


def _add_paragraph(doc, text, *, size=BODY_PT, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                   bold=False, underline=False, indent=None, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = LINE
    pf.space_before = SPACE
    pf.space_after = SPACE
    if indent is not None:
        pf.left_indent = indent
    for chunk, is_bold, is_italic in _inline(text):
        run = p.add_run(chunk)
        _style_run(run, size, bold=bold or is_bold, italic=is_italic, color=color)
        run.underline = underline
    return p


def _inline(text):
    """Quebra o texto em (trecho, negrito, itálico) lendo **negrito** e *itálico*."""
    out = []
    for part in re.split(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*))", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*"):
            out.append((part[1:-1], False, True))
        else:
            out.append((part, False, False))
    return out or [(text, False, False)]


def _add_prova(doc, numero, descricao):
    """Moldura vazia, com borda, onde o advogado cola o print do dossiê."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.text = ""

    titulo = cell.paragraphs[0]
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_before = Pt(6)
    titulo.paragraph_format.space_after = Pt(0)
    _style_run(titulo.add_run(f"[ {numero} - COLE O PRINT AQUI ]"), 11, bold=True,
               color=RGBColor(0xB0, 0x00, 0x00))

    legenda = cell.add_paragraph()
    legenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
    legenda.paragraph_format.space_before = Pt(0)
    legenda.paragraph_format.space_after = Pt(6)
    _style_run(legenda.add_run(descricao), 9, italic=True,
               color=RGBColor(0x55, 0x55, 0x55))

    _bordas(table, sz=8, val="dashed")
    doc.add_paragraph()


def _add_quadro(doc, linhas):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = ""
    for i, linha in enumerate(linhas):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        rotulo, _, valor = linha.partition(":")
        _style_run(p.add_run(f"{rotulo.strip()}: "), BODY_PT, bold=True)
        _style_run(p.add_run(valor.strip()), BODY_PT)
    _bordas(table, sz=8, val="single")
    doc.add_paragraph()


def _bordas(table, sz=8, val="single"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right"):
        el = borders.makeelement(qn(f"w:{edge}"), {})
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tbl_pr.append(borders)


def montar(texto, template):
    doc = Document(str(template))
    linhas = texto.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(linhas):
        linha = linhas[i].rstrip()
        i += 1
        if not linha.strip():
            continue

        if linha.strip() == "---":
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            continue

        if linha.strip() == "[[QUADRO]]":
            bloco = []
            while i < len(linhas) and linhas[i].strip() != "[[/QUADRO]]":
                if linhas[i].strip():
                    bloco.append(linhas[i].strip())
                i += 1
            i += 1
            _add_quadro(doc, bloco)
            continue

        m = re.match(r"\[\[\s*(PROVA[^|\]]*?)\s*\|\s*(.+?)\s*\]\]", linha.strip())
        if m:
            _add_prova(doc, m.group(1).upper(), m.group(2))
            continue

        if linha.startswith("## "):
            _add_paragraph(doc, linha[3:].strip(), bold=True)
            continue
        if linha.startswith("# "):
            _add_paragraph(doc, linha[2:].strip(), align=WD_ALIGN_PARAGRAPH.CENTER,
                           bold=True, underline=True)
            continue
        if linha.startswith("^ "):
            _add_paragraph(doc, linha[2:].strip(), align=WD_ALIGN_PARAGRAPH.CENTER)
            continue
        if linha.startswith(">> "):
            _add_paragraph(doc, linha[3:].strip(), size=EMENTA_PT, indent=INDENT)
            continue
        if linha.startswith("> "):
            _add_paragraph(doc, linha[2:].strip(), size=CITE_PT, indent=INDENT)
            continue

        _add_paragraph(doc, linha.strip())

    return doc


def main():
    ap = argparse.ArgumentParser(description="Monta a apelação em .docx no padrão BFAP.")
    ap.add_argument("--in", dest="entrada", required=True, help="arquivo .md com a peça")
    ap.add_argument("--out", dest="saida", required=True, help="arquivo .docx de saída")
    ap.add_argument("--template", default=None, help="modelo timbrado (default: assets/modelo-bfap.docx)")
    args = ap.parse_args()

    template = Path(args.template) if args.template else \
        Path(__file__).resolve().parent.parent / "assets" / "modelo-bfap.docx"
    if not template.exists():
        sys.exit(f"Template não encontrado: {template}")

    texto = Path(args.entrada).read_text(encoding="utf-8")
    doc = montar(texto, template)
    doc.save(args.saida)
    print(f"OK: {args.saida}")


if __name__ == "__main__":
    main()
