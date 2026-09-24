"""Gera as petições a partir da análise processual de cada caso (dados_djen.jsonl).

Cada linha do JSONL descreve um processo depois da leitura das publicações no DJEN:
    processo, gerar (bool), tipo, enderecamento, reu, autor, plural (bool), fato, ultimo_ato, orgao, obs

Uso:
    python gerar_de_dados.py dados_djen.jsonl MODELO.docx --pasta SAIDA [--data AAAA-MM-DD]

Minutas já existentes na pasta não são refeitas. Grava também a planilha de controle.
"""
import argparse
import datetime
import json
import os

import docx
from docx.oxml.ns import qn

from gerar_minutas import OAB_MARIO, _p, data_extenso, gravar_controle, tribunal

DURACAO = ("em observância ao princípio da razoável duração do processo (art. 5º, LXXVIII, da Constituição "
           "Federal) e ao dever de o juiz velar por ela (arts. 4º e 139, II, do Código de Processo Civil)")


def corpo_por_tipo(d):
    """Pedido (caixa-alta) e parágrafos do corpo, conforme o andamento real do processo."""
    quem = "as rés requerem" if d.get("plural") else "a ré requer"
    tipo = d["tipo"]
    if tipo == "remessa_recurso":
        return "a REMESSA DOS AUTOS À TURMA RECURSAL", [
            d["fato"],
            "Encerrado o processamento do recurso no juízo de origem, com o decurso do prazo de resposta previsto "
            "no art. 42, § 2º, da Lei nº 9.099/95, resta apenas o encaminhamento dos autos à instância revisora, "
            f"providência que independe de nova manifestação das partes. Assim, {DURACAO}, {quem} a remessa "
            "dos autos à Turma Recursal para o julgamento do recurso inominado.",
        ]
    if tipo == "transito_arquivamento":
        return "a CERTIFICAÇÃO DO TRÂNSITO EM JULGADO E O ARQUIVAMENTO DEFINITIVO DOS AUTOS", [
            d["fato"],
            "Nada mais havendo a ser decidido, e não subsistindo providência pendente a cargo das partes, "
            f"{quem}, {DURACAO}, a certificação do trânsito em julgado, caso ainda não lançada, e o arquivamento "
            "definitivo do feito, com as baixas de estilo.",
        ]
    if tipo == "julgamento_ed":
        return "o JULGAMENTO DOS EMBARGOS DE DECLARAÇÃO", [
            d["fato"],
            "Nos termos do art. 1.024, § 1º, do Código de Processo Civil, nos tribunais o relator apresentará os "
            "embargos de declaração em mesa na sessão subsequente. Assim, "
            f"{DURACAO}, {quem} a inclusão dos embargos de declaração em pauta e o seu julgamento.",
        ]
    if tipo == "sentenca":
        return "a CONCLUSÃO DOS AUTOS PARA PROLAÇÃO DE SENTENÇA", [
            d["fato"],
            "Os arts. 226 e 227 do Código de Processo Civil fixam prazo para a prolação de sentença, admitida sua "
            f"prorrogação apenas por motivo justificado. Assim, {DURACAO}, {quem} a imediata conclusão dos autos "
            "e a prolação de sentença.",
        ]
    if tipo == "prosseguimento":
        return "o REGULAR PROSSEGUIMENTO DO FEITO", [
            d["fato"],
            f"Desse modo, {DURACAO}, {quem} o prosseguimento do feito, com a conclusão dos autos para despacho "
            "ou, nada mais havendo a ser produzido, para sentença.",
        ]
    if tipo == "audiencia":
        return "a DESIGNAÇÃO DE AUDIÊNCIA OU O JULGAMENTO ANTECIPADO DO MÉRITO", [
            d["fato"],
            f"Assim, {DURACAO}, {quem} a designação de audiência de conciliação, instrução e julgamento ou, "
            "sendo desnecessária a produção de outras provas, o julgamento antecipado do mérito, nos termos do "
            "art. 355, I, do Código de Processo Civil.",
        ]
    raise ValueError(f"tipo desconhecido: {tipo}")


def montar(modelo, destino, d, pedido, corpo, oab, fecho):
    doc = docx.Document(modelo)
    body = doc.element.body
    tabela = body.find(qn("w:tbl"))
    for el in list(body):
        if el is tabela:
            break
        body.remove(el)
    plural = d.get("plural")
    qualif = ("já devidamente qualificadas e representadas" if plural else
              "já devidamente qualificada e representada")
    verbo = "vêm" if plural else "vem"
    indent = 2126
    novos = [
        _p([(d["enderecamento"], True)]),
        _p([]), _p([]),
        _p([(f"Autos nº {d['processo']}", True)]),
        _p([(d["reu"], True)]
           # Sem o nome da parte contrária numa fonte pública, a qualificação fica só "nos autos em epígrafe".
           + ([(f", {qualif} nos autos em epígrafe, em que contende{'m' if plural else ''} com ", False),
               (d["autor"], True)] if d.get("autor") and d["autor"] != "—" else
              [(f", {qualif} nos autos em epígrafe", False)])
           + [(f", {verbo}, respeitosamente, à presença de Vossa Excelência, por seus advogados infra-assinados, "
               "requerer ", False), (pedido, True), (", nos termos a seguir dispostos.", False)],
           esquerda=0, primeira_linha=0),
    ]
    novos += [_p([(c, False)], primeira_linha=indent) for c in corpo]
    novos += [
        _p([("Por fim, requer que as publicações ocorridas nestes autos sejam feitas exclusivamente em nome de ",
             False), ("MARIO THADEU LEME DE BARROS FILHO", True), (f" ({oab}), sob pena de nulidade.", False)],
           primeira_linha=indent),
        _p([]),
        _p([("Termos em que pede deferimento.", False)], alinhamento="center"),
        _p([(f"São Paulo, {fecho}.", False)], alinhamento="center"),
    ]
    for el in novos:
        tabela.addprevious(el)
    for cel in doc.tables[0]._cells:
        for par in cel.paragraphs:
            if par.text.strip().replace(" ", "") == "OAB/SP246.508":
                par.runs[0].text = oab
                for r in par.runs[1:]:
                    r.text = ""
    doc.save(destino)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dados")
    ap.add_argument("modelo")
    ap.add_argument("--pasta", required=True)
    ap.add_argument("--data")
    a = ap.parse_args()
    fecho = data_extenso(datetime.date.fromisoformat(a.data) if a.data else datetime.date.today())
    os.makedirs(a.pasta, exist_ok=True)

    controle = []
    with open(a.dados, encoding="utf-8") as f:
        registros = [json.loads(l) for l in f if l.strip()]
    for d in registros:
        uf, oab = OAB_MARIO.get(tribunal(d["processo"]), (None, None))
        linha = {"linha": "", "processo": d["processo"], "prioridade": "", "dias": "",
                 "providencia": d.get("tipo", "—"), "reu": d.get("reu", ""), "autor": d.get("autor", ""),
                 "fonte_autor": "DJEN", "oab": oab or "—",
                 "obs": f"Último ato: {d.get('ultimo_ato', '')}. Órgão: {d.get('orgao', '')}. {d.get('obs', '')}"}
        if not d.get("gerar"):
            controle.append({**linha, "status": "Sem minuta — pendência", "arquivo": ""})
            continue
        nome = f"{d['processo']} - petição de impulso.docx"
        destino = os.path.join(a.pasta, nome)
        if os.path.exists(destino):
            controle.append({**linha, "status": "Minuta já existia — não refeita", "arquivo": nome})
            continue
        pedido, corpo = corpo_por_tipo(d)
        montar(a.modelo, destino, d, pedido, corpo, oab, fecho)
        controle.append({**linha, "status": "Minuta pronta para revisão", "arquivo": nome})

    gravar_controle(os.path.join(a.pasta, "Controle - Peticionamento 30 dias.xlsx"), controle)
    prontas = sum(1 for c in controle if c["status"] == "Minuta pronta para revisão")
    print(f"{prontas} minutas geradas; {len(controle) - prontas} sem minuta/pendência")


if __name__ == "__main__":
    main()
