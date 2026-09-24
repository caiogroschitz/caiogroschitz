"""Filtra, na planilha de saneamento, os processos parados de um responsável e gera o lote.csv.

Uso:
    python lote.py PLANILHA.xlsx --responsavel "Caio Groschitz dos Santos Cruz" [--saida lote.csv]

O lote.csv é a lista de trabalho: as colunas `autor`, `enderecamento_astrea`, `obs_astrea`
e `status_astrea` ficam vazias para serem preenchidas com o que se lê no Astrea.
Se o lote.csv já existir, as linhas já preenchidas são preservadas (checkpoint).
"""
import argparse
import csv
import datetime
import os
import sys

import openpyxl

COLUNAS = [
    "linha_planilha", "processo", "cliente", "vara", "foro", "instancia", "data_ult_historico",
    "dias_parado", "prioridade", "providencia", "detalhamento", "etiquetas", "status_projuris",
    "autor", "enderecamento_astrea", "status_astrea", "obs_astrea",
]
CABECALHO_OBRIGATORIO = ["Nº do Processo", "Cliente", "Responsável (Astrea)", "Providência sugerida"]


def _data(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, (int, float)):
        return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))).isoformat()
    return (v or "").strip() if isinstance(v, str) else ""


def ler_planilha(caminho, responsavel):
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    ws = wb["Saneamento"] if "Saneamento" in wb.sheetnames else wb.worksheets[-1]
    linhas = list(ws.iter_rows(values_only=True))
    idx_cab = next(i for i, r in enumerate(linhas) if r and all(c in r for c in CABECALHO_OBRIGATORIO))
    cab = {str(c).strip(): i for i, c in enumerate(linhas[idx_cab]) if c}

    def g(r, nome):
        v = r[cab[nome]] if nome in cab and cab[nome] < len(r) else None
        return v.strip() if isinstance(v, str) else v

    responsaveis = set()
    saida = []
    for n, r in enumerate(linhas[idx_cab + 1:], start=idx_cab + 2):
        resp = g(r, "Responsável (Astrea)")
        if resp:
            responsaveis.add(resp)
        if resp != responsavel:
            continue
        if (g(r, "Peticionado?") or "").strip():
            continue  # já peticionado ou prejudicado
        saida.append({
            "linha_planilha": n,
            "processo": g(r, "Nº do Processo"),
            "cliente": g(r, "Cliente") or "",
            "vara": " ".join(str(g(r, "Vara") or "").split()),
            "foro": " ".join(str(g(r, "Foro") or "").split()),
            "instancia": str(g(r, "Inst.") or "").replace(".0", ""),
            "data_ult_historico": _data(g(r, "Data últ. histórico")),
            "dias_parado": int(float(g(r, "Dias parado") or 0)),
            "prioridade": g(r, "Prioridade") or "",
            "providencia": g(r, "Providência sugerida") or "",
            "detalhamento": g(r, "Detalhamento / pedido a ser formulado") or "",
            "etiquetas": g(r, "Etiquetas-chave (Astrea)") or "",
            "status_projuris": g(r, "Status (Projuris)") or "",
        })
    return saida, sorted(responsaveis)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("planilha")
    ap.add_argument("--responsavel", required=True)
    ap.add_argument("--saida", default="lote.csv")
    a = ap.parse_args()

    lote, responsaveis = ler_planilha(a.planilha, a.responsavel)
    if not lote:
        parecidos = [r for r in responsaveis if a.responsavel.split()[0].lower() in r.lower()]
        sys.exit(f"Nenhum processo pendente para '{a.responsavel}'. Nomes parecidos na planilha: {parecidos}")

    # Preserva o que já foi preenchido a partir do Astrea em execução anterior.
    anteriores = {}
    if os.path.exists(a.saida):
        with open(a.saida, newline="", encoding="utf-8") as f:
            anteriores = {r["processo"]: r for r in csv.DictReader(f)}
    for item in lote:
        ant = anteriores.get(item["processo"], {})
        for c in ("autor", "enderecamento_astrea", "status_astrea", "obs_astrea"):
            item[c] = ant.get(c, "")

    lote.sort(key=lambda x: -x["dias_parado"])
    with open(a.saida, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUNAS)
        w.writeheader()
        w.writerows(lote)
    print(f"{len(lote)} processos de {a.responsavel} gravados em {a.saida}")


if __name__ == "__main__":
    main()
